"""把 Markdown 报告排版为学术风格的 Word 文档。

处理本项目实际使用的语法子集：标题、段落、粗体、行内代码、表格、图片、
图注、有序/无序列表、参考文献编号条目。
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Times New Roman"
EA_FONT = "宋体"
MONO_FONT = "Consolas"


def set_run_font(run, size=10.5, bold=False, italic=False, mono=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = MONO_FONT if mono else BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_rich_text(paragraph, text: str, size=10.5, base_bold=False):
    """解析 **bold**、*italic*、`code` 并写入 run。"""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size, bold=base_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size - 0.5, bold=base_bold, mono=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size, bold=base_bold, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size, bold=base_bold)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r, row in enumerate(rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            text = row[c] if c < len(row) else ""
            add_rich_text(paragraph, text, size=9, base_bold=(r == 0))
    doc.add_paragraph()


def parse_table_block(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and re.match(r"^[\s:|-]+$", "|".join(rows[1])):
        rows.pop(1)
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="manuscript/MD6329_final_report.md")
    parser.add_argument("--output", default="manuscript/MD6329_final_report.docx")
    parser.add_argument("--figure-dir", default="manuscript")
    args = parser.parse_args()

    source = Path(args.input).read_text(encoding="utf-8")
    lines = source.split("\n")
    figure_root = Path(args.figure_dir)

    doc = Document()
    setup_document(doc)

    index = 0
    pending_caption = None
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if not stripped or stripped.startswith("---"):
            index += 1
            continue

        if stripped.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(10)
            run = paragraph.add_run(stripped[2:].strip())
            set_run_font(run, 16, bold=True)
            index += 1
            continue

        if stripped.startswith("## "):
            text = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", stripped[3:].strip())
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run(text)
            set_run_font(run, 12.5, bold=True)
            index += 1
            continue

        if stripped.startswith("### "):
            text = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", stripped[4:].strip())
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run(text)
            set_run_font(run, 11, bold=True, italic=True)
            index += 1
            continue

        if stripped.startswith("!["):
            match = re.match(r"!\[[^\]]*\]\(([^)]+)\)", stripped)
            index += 1
            caption = ""
            if index < len(lines) and lines[index].strip().startswith("**Figure"):
                caption = lines[index].strip()
                caption = re.sub(r"^\*\*Figure\s*[0-9]+\.\*\*\s*", "", caption)
                index += 1
            if match:
                image_path = Path(match.group(1))
                if not image_path.is_absolute():
                    image_path = figure_root / image_path
                if image_path.exists():
                    picture = doc.add_paragraph()
                    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    picture.add_run().add_picture(str(image_path), width=Inches(6.2))
            if caption:
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(10)
                run = paragraph.add_run(caption)
                set_run_font(run, 9)
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[index + 1].strip()):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(doc, parse_table_block(block))
            continue

        if stripped.startswith("**Table"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = True
            add_rich_text(paragraph, stripped, size=9.5)
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.space_after = Pt(6)
            add_rich_text(paragraph, stripped[2:].strip(), size=9.5)
            index += 1
            continue

        if re.match(r"^[-*]\s", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            add_rich_text(paragraph, stripped[2:].strip())
            index += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(2)
            add_rich_text(paragraph, stripped[3:].strip())
            index += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Inches(0.25)
        add_rich_text(paragraph, stripped)
        index += 1

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
